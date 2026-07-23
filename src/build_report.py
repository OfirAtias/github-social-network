# -*- coding: utf-8 -*-
"""Create the editable eight-page academic report from fresh output tables."""

from pathlib import Path
import sys

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SRC_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SRC_DIR.parent
sys.path.insert(0, str(SRC_DIR))

from final_validation import collect_key_results


REPORT_DIR = PROJECT_DIR / "report"
REPORT_PATH = REPORT_DIR / "report_source.docx"
TABLES_DIR = PROJECT_DIR / "outputs" / "tables"
FIGURES_DIR = PROJECT_DIR / "outputs" / "figures"

# Compact-reference-guide base with a named academic page-limit override.
PAGE_WIDTH_DXA = 12_240
PAGE_HEIGHT_DXA = 15_840
MARGIN_DXA = 900
CONTENT_WIDTH_DXA = PAGE_WIDTH_DXA - 2 * MARGIN_DXA
TABLE_INDENT_DXA = 100
CELL_MARGIN_DXA = {"top": 55, "bottom": 55, "start": 100, "end": 100}
INK = "172B4D"
ACCENT = "315C8D"
ACCENT_2 = "D57837"
MUTED = "5F6B7A"
LIGHT_FILL = "E8EEF5"
PALE_FILL = "F4F6F9"


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _set_run_font(
    run,
    *,
    size: float,
    bold: bool = False,
    color: str = "000000",
    italic: bool = False,
) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)


def _set_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    paragraph.alignment = alignment
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    normal.font.size = Pt(9.3)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.03

    heading_tokens = {
        "Heading 1": (13.5, 7, 3, ACCENT),
        "Heading 2": (11.0, 5, 2, ACCENT),
        "Heading 3": (10.0, 4, 1.5, INK),
    }
    for style_name, (size, before, after, color) in heading_tokens.items():
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = _rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def _configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    margin = Inches(MARGIN_DXA / 1440)
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)


def _add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    _set_run_font(run, size=8, color=MUTED)


def _set_header_footer(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    _set_rtl(paragraph, WD_ALIGN_PARAGRAPH.RIGHT)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("תחום התמחות ברשת GitHub | פרויקט גמר")
    _set_run_font(run, size=7.8, color=MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    _set_rtl(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run("עמוד ")
    _set_run_font(run, size=8, color=MUTED)
    _add_field(paragraph, "PAGE")
    run = paragraph.add_run(" מתוך 8")
    _set_run_font(run, size=8, color=MUTED)


def _add_paragraph(
    document: Document,
    text: str,
    *,
    bold_lead: str | None = None,
    size: float = 9.3,
    color: str = "000000",
    after: float = 3,
    align=WD_ALIGN_PARAGRAPH.RIGHT,
    keep_with_next: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    _set_rtl(paragraph, align)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        _set_run_font(lead, size=size, bold=True, color=color)
        tail = paragraph.add_run(text[len(bold_lead) :])
        _set_run_font(tail, size=size, color=color)
    else:
        run = paragraph.add_run(text)
        _set_run_font(run, size=size, color=color)


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    _set_rtl(paragraph)
    run = paragraph.add_run(text)
    size = {1: 13.5, 2: 11.0, 3: 10.0}[level]
    color = ACCENT if level < 3 else INK
    _set_run_font(run, size=size, bold=True, color=color)


def _add_bullet(document: Document, text: str) -> None:
    """Add a real Word-numbering bullet with compact report spacing."""
    paragraph = document.add_paragraph(style="List Bullet")
    _set_rtl(paragraph)
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.14)
    paragraph.paragraph_format.space_after = Pt(1.2)
    run = paragraph.add_run(text)
    _set_run_font(run, size=8.8)


def _shade_cell(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def _set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in CELL_MARGIN_DXA.items():
        tag = qn(f"w:{edge}")
        node = tc_mar.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _set_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError("Table widths must equal the configured content width.")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        _prevent_row_split(row)
        for cell, width in zip(row.cells, widths_dxa):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def _set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    size: float = 8.2,
    color: str = "000000",
    alignment=WD_ALIGN_PARAGRAPH.RIGHT,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    _set_rtl(paragraph, alignment)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    _set_run_font(run, size=size, bold=bold, color=color)


def _add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
    *,
    font_size: float = 8.2,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        _shade_cell(cell, LIGHT_FILL)
        _set_cell_text(cell, header, bold=True, size=font_size, color=INK)
    _set_repeat_header(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            _set_cell_text(cell, value, size=font_size)
    _set_table_geometry(table, widths_dxa)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def _add_picture(
    document: Document,
    filename: str,
    caption: str,
    *,
    width: float,
) -> None:
    path = FIGURES_DIR / filename
    if not path.is_file():
        raise FileNotFoundError(path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    caption_p = document.add_paragraph()
    _set_rtl(caption_p, WD_ALIGN_PARAGRAPH.CENTER)
    caption_p.paragraph_format.space_before = Pt(1)
    caption_p.paragraph_format.space_after = Pt(3)
    caption_run = caption_p.add_run(caption)
    _set_run_font(caption_run, size=7.8, italic=True, color=MUTED)


def _page_break(document: Document) -> None:
    document.add_page_break()


def _r(value: float, digits: int = 3) -> str:
    return f"{float(value):.{digits}f}"


def main() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    key = collect_key_results(PROJECT_DIR)
    basic = pd.read_csv(TABLES_DIR / "network_basic_stats.csv").iloc[0]
    centrality = pd.read_csv(TABLES_DIR / "centrality_top_developers.csv")
    overlap = pd.read_csv(TABLES_DIR / "centrality_top_overlap.csv")
    community_metrics = pd.read_csv(TABLES_DIR / "community_metrics.csv").iloc[0]
    communities = pd.read_csv(TABLES_DIR / "community_summary.csv")
    classification = pd.read_csv(TABLES_DIR / "classification_summary.csv")
    thresholds = pd.read_csv(TABLES_DIR / "threshold_tuning_summary.csv").set_index(
        "evaluation_type"
    )
    ablation = pd.read_csv(TABLES_DIR / "ablation_model_results_summary.csv")
    embeddings = pd.read_csv(TABLES_DIR / "embedding_classification_summary.csv")
    reconstruction = pd.read_csv(TABLES_DIR / "edge_reconstruction_summary.csv")
    reconstruction_groups = pd.read_csv(
        TABLES_DIR / "edge_reconstruction_same_vs_cross_summary.csv"
    )
    null_homophily = pd.read_csv(TABLES_DIR / "null_model_homophily.csv").set_index(
        "null_model"
    )
    null_modularity = pd.read_csv(TABLES_DIR / "null_model_modularity.csv").iloc[0]
    network_models = pd.read_csv(TABLES_DIR / "network_vs_random_models.csv")
    powerlaw = pd.read_csv(TABLES_DIR / "powerlaw_fit.csv").iloc[0]

    document = Document()
    _configure_styles(document)
    section = document.sections[0]
    _configure_section(section)
    _set_header_footer(section)

    # Page 1: integrated title, abstract and research questions.
    title = document.add_paragraph()
    _set_rtl(title, WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("תחום התמחות של מפתחי GitHub מתוך רשת מעקב הדדית")
    _set_run_font(run, size=18, bold=True, color=INK)
    subtitle = document.add_paragraph()
    _set_rtl(subtitle, WD_ALIGN_PARAGRAPH.CENTER)
    subtitle.paragraph_format.space_after = Pt(2)
    run = subtitle.add_run("פרויקט גמר — אלגוריתמים ברשתות")
    _set_run_font(run, size=10.5, bold=True, color=ACCENT)
    meta = document.add_paragraph()
    _set_rtl(meta, WD_ALIGN_PARAGRAPH.CENTER)
    meta.paragraph_format.space_after = Pt(5)
    run = meta.add_run(
        "נועה ירין לוי (208101196) · ג'ימי חלאוי (208301176) · "
        "אופיר אטיאס (211446935) | מרצה: ד\"ר אבנר פריאל"
    )
    _set_run_font(run, size=8.8, color=MUTED)

    _add_heading(document, "תקציר", 1)
    _add_paragraph(
        document,
        f"העבודה מנתחת רשת לא־מכוונת של {key['network_nodes']:,} מפתחי GitHub "
        f"ו־{key['network_edges']:,} קשרי מעקב הדדיים, ובוחנת ארבע שאלות: "
        "סיווג תחום מקצועי מתוך מבנה הרשת, ערך הרשת מול תוכן, שחזור קשתות "
        "שהוסתרו ומובהקות מבנית מול מודלי אפס. כל הניסויים חוזרים בחמש חלוקות "
        "קבועות; תוויות הבדיקה אינן משמשות בחישוב מאפיינים או בכיול. שיעור "
        f"הקשתות התוך־תחומיות הוא {_r(100 * key['homophily_rate'], 1)}%. "
        f"המודל החזק בסיווג הידני הגיע ל־AUC {_r(key['best_classification_roc_auc'])}; "
        f"DeepWalk ו־Node2Vec הגיעו ל־{_r(key['deepwalk_roc_auc'])} "
        f"ו־{_r(key['node2vec_roc_auc'])}. בשחזור קשתות, "
        f"{key['best_reconstruction_heuristic']} הוביל עם "
        f"AUC {_r(key['best_reconstruction_roc_auc'])}. "
        "התוצאות מתוארות כקשרים חיזויים בתנאי הניסוי ואינן מוכיחות סיבתיות."
    )
    _add_heading(document, "שאלות המחקר", 1)
    questions = [
        "RQ1: באיזו מידה שכנות, מרכזיות, קהילות וייצוגים נלמדים מאפשרים לסווג את תחום המפתח ללא תוויות test?",
        "RQ2: מהו הערך השולי של מבנה הרשת מול מאפייני תוכן, ומה מוסיף שילובם?",
        "RQ3: האם ניתן לשחזר קשתות מעקב שהוסתרו אקראית, והאם קשתות תוך־תחומיות קלות יותר לשחזור מקשתות בין־תחומיות?",
        "RQ4: האם ההומופיליה והמודולריות חריגות מול מודלי אפס, ואילו מאפיינים משחזרים מודלי ER, BA, Small-World ו־Configuration?",
    ]
    for question in questions:
        _add_paragraph(document, question, size=8.9, after=1.5)
    _add_heading(document, "הנתונים והיקף הניתוח", 1)
    _add_table(
        document,
        ["צמתים", "קשתות", "דרגה ממוצעת", "צפיפות", "רכיבי קשירות", "תוויות חסרות"],
        [[
            f"{int(basic['num_nodes']):,}",
            f"{int(basic['num_edges']):,}",
            _r(basic["average_degree"]),
            _r(basic["density"], 6),
            str(int(basic["num_connected_components"])),
            "0",
        ]],
        [1500, 1500, 1800, 1600, 1800, 2240],
        font_size=8.0,
    )

    # Page 2: methods and the exact ten logical classification models.
    _page_break(document)
    _add_heading(document, "1. מסגרת מתודולוגית ושיטות", 1)
    _add_paragraph(
        document,
        "בסיווג הצמתים מבנה הגרף גלוי לכל הצמתים (transductive), אך תוויות test "
        "מוסתרות. מאפייני שכנים וקהילה מחושבים מתוויות train בלבד. בכל seed "
        "מתבצעת חלוקה מרובדת 70/30; כיול סף משתמש ב־validation פנימי מתוך train.",
    )
    _add_heading(document, "ניתוח הרשת", 2)
    _add_paragraph(
        document,
        "נמדדו Degree Centrality, Clustering Coefficient, PageRank, Betweenness "
        "בקירוב, HITS Hub יחיד, Louvain ומודולריות. התפלגות הדרגות נבחנה בזהירות "
        "באמצעות התאמת Power Law והשוואה ללוג־נורמלית ולמעריכית. הרשת הושוותה "
        "ל־ER, BA, Watts-Strogatz ולגרף משמר־דרגות.",
    )
    _add_heading(document, "מודלי בסיס ופיצ'רים ידניים", 2)
    _add_table(
        document,
        ["מודל", "קלט", "מטרת ההשוואה"],
        [
            ["Majority Baseline", "שכיחות train", "קו בסיס ללא מבנה"],
            ["Neighbor Majority", "תוויות שכני train", "אות שכונתי מקומי"],
            ["LogReg Graph", "מרכזיות ומבנה", "פיצ'רים ידניים"],
            ["LogReg Graph + Community", "מבנה וקהילה", "ערך קהילה מעבר למבנה"],
        ],
        [3200, 3200, 4040],
        font_size=8.1,
    )
    _add_heading(document, "ייצוגים נלמדים ומודלי תוכן", 2)
    _add_table(
        document,
        ["מודל", "קלט", "תפקיד"],
        [
            ["LogReg Content", "מאפייני MUSAE", "נקודת ייחוס ללא רשת"],
            ["LogReg Hybrid", "רשת + תוכן", "מידע משלים"],
            ["LogReg DeepWalk", "הילוכים אחידים", "ייצוג מבני נלמד"],
            ["DeepWalk + ידני", "הטמעה + מבנה/קהילה", "שילוב ייצוגים"],
            ["LogReg Node2Vec", "הילוכים מוטים", "ייצוג מבני נלמד"],
            ["Node2Vec + ידני", "הטמעה + מבנה/קהילה", "שילוב ייצוגים"],
        ],
        [3200, 3200, 4040],
        font_size=7.9,
    )
    _add_heading(document, "שחזור קשתות ומודלי אפס", 2)
    _add_paragraph(
        document,
        "בשחזור קשתות מוגן עץ פורש וקשתות test נדגמות רק מן הקשתות האחרות; "
        "כך גרף האימון נשאר קשיר. non-edges מותאמים לפי Web-Web, ML-ML ו־Web-ML. "
        "נבדקים Common Neighbors, Jaccard, Adamic-Adar, Preferential Attachment, "
        "Same Community, DeepWalk ו־Node2Vec. ערכי p אמפיריים מחושבים עם "
        "(b+1)/(B+1), ולכן אינם אפס.",
    )

    # Page 3: structural analysis and communities.
    _page_break(document)
    _add_heading(document, "2. מבנה הרשת, מרכזיות וקהילות", 1)
    _add_paragraph(
        document,
        f"הרשת קשירה, צפופה מעט בלבד ({basic['density']:.6f}) ובעלת דרגת שיא "
        f"{int(pd.read_csv(TABLES_DIR / 'network_vs_random_models.csv').iloc[0]['max_degree']):,}. "
        "Degree ו־PageRank מדגישים מוקדי קשרים, בעוד Betweenness מדגישה תיווך "
        "בין אזורים; לכן טבלאות המובילים אינן זהות.",
    )
    top_rows = []
    for metric in ("Degree Centrality", "PageRank", "Betweenness Centrality"):
        row = centrality[centrality["metric"].str.startswith(metric)].sort_values("rank").iloc[0]
        top_rows.append(
            [
                metric.replace(" Centrality", ""),
                str(row["developer_name"]),
                str(int(row["node_id"])),
                f"{int(row['degree']):,}",
                f"{float(row['score']):.6f}",
            ]
        )
    _add_table(
        document,
        ["מדד", "מפתח מוביל", "Node ID", "דרגה", "ציון"],
        top_rows,
        [1800, 3000, 1500, 1300, 2840],
        font_size=8.0,
    )
    degree_pr = overlap[
        (
            overlap["metric_a"].str.startswith("Degree Centrality")
            & overlap["metric_b"].str.startswith("PageRank")
        )
        | (
            overlap["metric_b"].str.startswith("Degree Centrality")
            & overlap["metric_a"].str.startswith("PageRank")
        )
    ]["jaccard_overlap"].iloc[0]
    degree_bet = overlap[
        (
            overlap["metric_a"].str.startswith("Degree Centrality")
            & overlap["metric_b"].str.startswith("Betweenness")
        )
        | (
            overlap["metric_b"].str.startswith("Degree Centrality")
            & overlap["metric_a"].str.startswith("Betweenness")
        )
    ]["jaccard_overlap"].iloc[0]
    _add_paragraph(
        document,
        f"חפיפת Jaccard בין 20 המובילים ב־Degree וב־PageRank היא "
        f"{degree_pr:.3f}, לעומת {degree_bet:.3f} בין Degree וב־Betweenness. "
        "הפער ממחיש שמתווכים אינם בהכרח הצמתים בעלי מספר הקשרים הגדול ביותר.",
    )
    _add_picture(
        document,
        "community_size_and_composition.png",
        "איור 1. שתים־עשרה קהילות Louvain הגדולות: גודל והרכב תחומי.",
        width=6.15,
    )
    _add_paragraph(
        document,
        f"Louvain מצא {int(community_metrics['number_of_communities'])} קהילות "
        f"עם מודולריות {community_metrics['modularity']:.3f}. הקהילה הגדולה "
        f"מכילה {int(community_metrics['largest_community_size']):,} צמתים; "
        f"הטוהר הממוצע המשוקלל הוא {community_metrics['size_weighted_mean_purity']:.3f}. "
        "הרכבים שונים בין קהילות מצביעים על קשר בין מבנה לתחום, אך קהילות "
        "מעורבות מונעות פירוש דטרמיניסטי.",
    )

    # Page 4: classification.
    _page_break(document)
    _add_heading(document, "3. RQ1 — סיווג צמתים", 1)
    class_rows = []
    for _, row in classification.sort_values("roc_auc_mean").iterrows():
        class_rows.append(
            [
                str(row["model"]).replace("Logistic Regression - ", "LogReg "),
                _r(row["macro_f1_mean"]),
                _r(row["macro_f1_std"]),
                _r(row["roc_auc_mean"]),
                _r(row["roc_auc_std"]),
            ]
        )
    _add_table(
        document,
        ["מודל", "Macro F1", "סטיית תקן", "ROC-AUC", "סטיית תקן"],
        class_rows,
        [4100, 1600, 1500, 1600, 1640],
        font_size=8.0,
    )
    _add_picture(
        document,
        "classification_comparison.png",
        "איור 2. ביצועי סיווג הצמתים על פני חמש חלוקות; פסי השגיאה הם סטיית תקן.",
        width=6.25,
    )
    default_threshold = thresholds.loc["default_threshold_0.5"]
    tuned_threshold = thresholds.loc["tuned_threshold_validation_f1"]
    _add_paragraph(
        document,
        f"הוספת קהילה העלתה את ROC-AUC ל־"
        f"{key['best_classification_roc_auc']:.3f}. סף שנבחר על validation בלבד "
        f"היה {tuned_threshold['selected_threshold_mean']:.3f} בממוצע; הוא שינה "
        f"את Precision מ־{default_threshold['precision_mean']:.3f} ל־"
        f"{tuned_threshold['precision_mean']:.3f} ואת Recall מ־"
        f"{default_threshold['recall_mean']:.3f} ל־{tuned_threshold['recall_mean']:.3f}. "
        "השינוי הוא trade-off תפעולי ואינו שיפור מובטח בכל מדד.",
    )
    _add_paragraph(
        document,
        "מסקנת RQ1: מבנה הרשת מספק אות חיזויי מעל קווי הבסיס; קהילה מוסיפה "
        "מידע מעבר לשכנות ולמרכזיות. המסקנה מוגבלת לפרוטוקול transductive ולתווית "
        "הבינארית שנבדקה.",
        bold_lead="מסקנת RQ1:",
    )

    # Page 5: ablation and embeddings.
    _page_break(document)
    _add_heading(document, "4. RQ2 — רשת, תוכן וייצוגים נלמדים", 1)
    comparison_rows = []
    for _, row in ablation.iterrows():
        comparison_rows.append(
            [
                str(row["model"]).replace("Logistic Regression - ", "LogReg "),
                _r(row["macro_f1_mean"]),
                _r(row["roc_auc_mean"]),
                _r(row["roc_auc_std"]),
            ]
        )
    for _, row in embeddings.iterrows():
        comparison_rows.append(
            [
                str(row["model"]).replace("Logistic Regression - ", "LogReg "),
                _r(row["macro_f1_mean"]),
                _r(row["roc_auc_mean"]),
                _r(row["roc_auc_std"]),
            ]
        )
    _add_table(
        document,
        ["ייצוג/מקור מידע", "Macro F1", "ROC-AUC", "סטיית תקן AUC"],
        comparison_rows,
        [5600, 1600, 1600, 1640],
        font_size=7.8,
    )
    _add_picture(
        document,
        "representation_comparison.png",
        "איור 3. השוואה הוגנת בין פיצ'רים ידניים, תוכן, DeepWalk, Node2Vec והיברידי.",
        width=6.3,
    )
    hybrid = ablation[ablation["model"].str.contains("Hybrid", regex=False)].iloc[0]
    content = ablation[ablation["model"].str.contains("Content Features Only", regex=False)].iloc[0]
    network_only = ablation[ablation["model"].str.contains("Network Only", regex=False)].iloc[0]
    _add_paragraph(
        document,
        f"תוכן לבדו הגיע ל־AUC {content['roc_auc_mean']:.3f}, מול "
        f"{network_only['roc_auc_mean']:.3f} בפיצ'רים ידניים של הרשת ו־"
        f"{hybrid['roc_auc_mean']:.3f} בשילוב. DeepWalk לבדו הגיע ל־"
        f"{key['deepwalk_roc_auc']:.3f} ו־Node2Vec ל־{key['node2vec_roc_auc']:.3f}. "
        "הדמיון בביצועים מצביע על חפיפה חיזויית בתנאי הניסוי; הוא אינו מוכיח "
        "שקילות בין מקורות המידע.",
    )
    _add_paragraph(
        document,
        "מסקנת RQ2: פיצ'רי תוכן חזקים מן הייצוג הידני של הרשת, אך השילוב משפר "
        "את התוצאה. הטמעות הילוכים מרכזות מידע מבני שימושי ומצמצמות חלק מן הפער.",
        bold_lead="מסקנת RQ2:",
    )

    # Page 6: edge reconstruction.
    _page_break(document)
    _add_heading(document, "5. RQ3 — שחזור קשתות קשיר", 1)
    _add_paragraph(
        document,
        "בכל seed מוגנות כל קשתות עץ פורש, ולכן הסרת 10% מהקשתות שאינן בעץ "
        "אינה מפרקת את גרף האימון. בכל חמש החלוקות התקבל רכיב קשיר אחד וללא "
        "צמתים מבודדים. שמירה זו נדרשת במיוחד למסלולים, Louvain והטמעות המבוססות "
        "הילוכים. לכל positive הותאם non-edge מאותו סוג קצוות.",
    )
    reconstruction_rows = []
    for _, row in reconstruction.sort_values("roc_auc_mean", ascending=False).iterrows():
        reconstruction_rows.append(
            [
                str(row["heuristic"]),
                _r(row["roc_auc_mean"]),
                _r(row["roc_auc_std"]),
                _r(row["average_precision_mean"]),
                _r(row["average_precision_std"]),
            ]
        )
    _add_table(
        document,
        ["שיטה", "ROC-AUC", "סטיית תקן", "Average Precision", "סטיית תקן"],
        reconstruction_rows,
        [3700, 1500, 1400, 2100, 1740],
        font_size=7.7,
    )
    _add_picture(
        document,
        "edge_reconstruction_same_vs_cross.png",
        "איור 4. השוואה ישירה same-field מול cross-field תחת דגימה מאוזנת.",
        width=6.35,
    )
    best_group = reconstruction_groups[
        reconstruction_groups["heuristic"] == key["best_reconstruction_heuristic"]
    ].set_index("edge_group")
    same_auc = best_group.loc["same-field", "roc_auc_mean"]
    cross_auc = best_group.loc["cross-field", "roc_auc_mean"]
    same_ap = best_group.loc["same-field", "average_precision_mean"]
    cross_ap = best_group.loc["cross-field", "average_precision_mean"]
    direction = "גבוה" if same_auc > cross_auc else "נמוך"
    _add_paragraph(
        document,
        f"{key['best_reconstruction_heuristic']} הוביל עם AUC "
        f"{key['best_reconstruction_roc_auc']:.3f}. באותה שיטה, same-field קיבל "
        f"AUC {same_auc:.3f} ו־AP {same_ap:.3f}, לעומת AUC {cross_auc:.3f} "
        f"ו־AP {cross_ap:.3f} ב־cross-field. לכן השחזור התוך־תחומי {direction} "
        "יותר בפרוטוקול זה. התוצאה מתארת יכולת שחזור ואינה מוכיחה מנגנון סיבתי.",
    )

    # Page 7: null models and network-model comparison.
    _page_break(document)
    _add_heading(document, "6. RQ4 — מודלי אפס ומודלי רשת", 1)
    _add_table(
        document,
        ["מודל אפס", "נצפה", "ממוצע אפס", "סטיית תקן", "p אמפירי"],
        [
            [
                "ערבוב תוויות",
                _r(null_homophily.loc["label_shuffle", "observed"]),
                _r(null_homophily.loc["label_shuffle", "null_mean"]),
                _r(null_homophily.loc["label_shuffle", "null_std"]),
                _r(null_homophily.loc["label_shuffle", "empirical_p"], 6),
            ],
            [
                "גרף משמר־דרגות",
                _r(null_homophily.loc["degree_preserving_rewire", "observed"]),
                _r(null_homophily.loc["degree_preserving_rewire", "null_mean"]),
                _r(null_homophily.loc["degree_preserving_rewire", "null_std"]),
                _r(null_homophily.loc["degree_preserving_rewire", "empirical_p"], 6),
            ],
        ],
        [2800, 1600, 2000, 1800, 2240],
        font_size=8.0,
    )
    model_rows = []
    for _, row in network_models.iterrows():
        model_rows.append(
            [
                str(row["network"]),
                f"{int(row['edges']):,}",
                _r(row["avg_clustering"], 4),
                _r(row["avg_shortest_path_sampled"]),
                str(int(row["max_degree"])),
            ]
        )
    _add_table(
        document,
        ["רשת", "קשתות", "Clustering", "מרחק ממוצע", "דרגת שיא"],
        model_rows,
        [4000, 1800, 1600, 1800, 1240],
        font_size=7.5,
    )
    _add_picture(
        document,
        "homophily_null_distribution.png",
        "איור 5. הומופיליה נצפית מול ערבוב תוויות וגרפים משמרי־דרגות.",
        width=5.9,
    )
    _add_paragraph(
        document,
        f"המודולריות הנצפית היא {null_modularity['observed_modularity']:.3f}, "
        f"מול {null_modularity['rewired_null_mean']:.3f} בממוצע בגרפים "
        f"משמרי־דרגות (p={null_modularity['empirical_p']:.6f}). Power Law נותן "
        f"α={powerlaw['alpha']:.2f}; ההשוואה ללוג־נורמלית אינה מספקת בסיס "
        "להכרזה גורפת על scale-free, ולכן המסקנה הזהירה היא זנב כבד. ER אינו "
        "משחזר התקבצות או דרגת שיא; BA משחזר חלק מן הזנב; Small-World משחזר "
        "התקבצות ומרחקים קצרים; הגרף משמר־הדרגות מבודד את תרומת רצף הדרגות.",
    )

    # Page 8: synthesis, limitations, conclusions and compact references.
    _page_break(document)
    _add_heading(document, "7. דיון מסכם ומגבלות", 1)
    _add_paragraph(
        document,
        "התמונה המצטברת עקבית בין ארבע השאלות: יש קשר תיאורי חזק בין מבנה "
        "הרשת לתחום, קהילות והטמעות מוסיפות כוח חיזויי, ותוכן מספק אות חזק "
        "נוסף. מודלי האפס מראים שרצף הדרגות מסביר חלק מן ההומופיליה אך לא את "
        "כולה. בשחזור הקשתות, הצלחת Preferential Attachment עקבית עם חשיבות "
        "הדרגות, אך אינה מוכיחה שהעדפת צמתים בעלי דרגה גבוהה היא מנגנון יצירת הקשר.",
    )
    _add_heading(document, "מגבלות", 2)
    limitations = [
        "הנתונים הם תמונת מצב מיוני 2019 ואינם מייצגים בהכרח את GitHub כיום.",
        "תווית היעד בינארית ונגזרה ממידע פרופיל; ייתכן רעש וחסר בתחום המקצועי.",
        "הסיווג transductive, ולכן אינו מדמה צומת חדש ללא קשרים ידועים.",
        "Betweenness מחושב בקירוב, והטמעות נבדקו בתקציב פרמטרים קבוע ללא חיפוש רחב.",
        "שחזור קשתות שהוסתרו אינו תחזית זמן ואינו מזהה סיבתיות או כיווניות.",
        "מודלי האפס והמודלים הגנרטיביים מבודדים מאפיינים מסוימים ואינם תיאור מלא של תהליך הרשת.",
    ]
    for item in limitations:
        _add_bullet(document, item)
    _add_heading(document, "מסקנות", 2)
    _add_paragraph(
        document,
        "RQ1 נענה בחיוב מוגבל: מבנה הרשת מאפשר סיווג מעל קווי הבסיס. RQ2 מראה "
        "חפיפה חלקית ומידע משלים בין רשת לתוכן. RQ3 מראה ששחזור קשתות אפשרי "
        "תוך שמירת קשירות, והפער same/cross תלוי בשיטה ובפרוטוקול. RQ4 מראה "
        "שההומופיליה והמודולריות חריגות ביחס למודלי האפס שנבדקו. הממצאים "
        "תומכים בערך האנליטי של רשתות מקצועיות, תוך שמירה על ניסוח תיאורי וזהיר.",
    )
    _add_heading(document, "מקורות", 2)
    references = [
        "Rozemberczki, Allen & Sarkar (2019), MUSAE.",
        "Blondel et al. (2008), Fast unfolding of communities.",
        "Perozzi, Al-Rfou & Skiena (2014), DeepWalk.",
        "Grover & Leskovec (2016), node2vec.",
        "Liben-Nowell & Kleinberg (2003), Link Prediction.",
        "Clauset, Shalizi & Newman (2009), Power-law distributions.",
        "Page et al. (1999), PageRank.",
        "מצגות ותרגולי הקורס: מרכזיות, קהילות, מודלי רשת, שחזור קשתות והטמעות.",
    ]
    for index, reference in enumerate(references, start=1):
        _add_paragraph(document, f"{index}. {reference}", size=7.8, after=0.6)

    # Ensure all sections (including any Word-created continuation sections)
    # use the same compact page geometry and furniture.
    for section in document.sections:
        _configure_section(section)
        _set_header_footer(section)

    document.core_properties.title = (
        "תחום התמחות של מפתחי GitHub מתוך רשת מעקב הדדית"
    )
    document.core_properties.subject = "Final project - network algorithms"
    document.core_properties.author = (
        "Noa Yarin Levy; Jimi Halawi; Ofir Atias"
    )
    document.save(REPORT_PATH)
    print(f"Built report source: {REPORT_PATH}")


if __name__ == "__main__":
    main()
